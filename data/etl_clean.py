"""
Step 2 — ETL 清洗
Input : raw_data/shipping_2024_raw.xlsx
Output: output/cleaned_records.json

清洗項目：
  - 日期：辨識 4 種格式 + 遺漏值處理
  - 車種：6 種寫法 → HGV / LGV / LGV-S / UNKNOWN
  - 重量：kg / 噸 / 帶單位文字 / 遺漏值 → 統一 kg
  - 資料品質旗標 per record
"""

import json
import os
import openpyxl
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).parent

# ── 車種對應表 ─────────────────────────────────────────────────────────────────

VEHICLE_MAP = {
    "大型貨車": "HGV", "大型貨车": "HGV", "大貨車": "HGV",
    "HGV": "HGV",     "大型": "HGV",    "10噸車": "HGV",
    "中型貨車": "LGV", "中型货车": "LGV", "3.5t-7.5t貨車": "LGV",
    "LGV": "LGV",     "中型": "LGV",    "5噸車": "LGV",
    "小型貨車": "LGV-S", "小貨車": "LGV-S", "1噸車": "LGV-S", "LGV-S": "LGV-S",
}

VEHICLE_LABELS = {
    "HGV":     "大型貨車 (HGV >7.5t)",
    "LGV":     "中型貨車 (LGV 3.5-7.5t)",
    "LGV-S":   "小型貨車 (LGV <3.5t)",
    "UNKNOWN": "未知車種",
}

# Control Tower 的 vehicle 字串
VEHICLE_CT = {
    "HGV": "large", "LGV": "medium", "LGV-S": "small", "UNKNOWN": "unknown",
}

# ── 清洗函式 ──────────────────────────────────────────────────────────────────

DATE_FORMATS = ["%Y/%m/%d", "%Y-%m-%d", "%m/%d/%Y", "%Y.%m.%d"]

def parse_date(raw) -> tuple[str | None, str]:
    if raw is None or str(raw).strip() == "":
        return None, "WARN_DATE"
    raw = str(raw).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date().isoformat(), "OK"
        except ValueError:
            continue
    return None, "WARN_DATE"


def normalize_vehicle(raw) -> tuple[str, str]:
    if raw is None or str(raw).strip() == "":
        return "UNKNOWN", "WARN_VEHICLE"
    raw = str(raw).strip()
    result = VEHICLE_MAP.get(raw)
    if result:
        return result, "OK"
    for key, val in VEHICLE_MAP.items():
        if key.lower() == raw.lower():
            return val, "OK"
    return "UNKNOWN", "WARN_VEHICLE"


def normalize_weight(raw) -> tuple[int | None, str]:
    if raw is None or str(raw).strip() == "":
        return None, "WARN_WEIGHT"
    raw = str(raw).strip()

    if "噸" in raw:
        try:
            return round(float(raw.replace("噸", "").strip()) * 1000), "OK"
        except ValueError:
            return None, "WARN_WEIGHT"

    if "kg" in raw.lower():
        try:
            return round(float(raw.lower().replace("kg", "").strip())), "OK"
        except ValueError:
            return None, "WARN_WEIGHT"

    try:
        val = float(raw)
        if val <= 0:
            return None, "WARN_WEIGHT"
        if val < 30:
            # Ambiguous: likely in tonnes (no unit label)
            return round(val * 1000), "WARN_WEIGHT"
        return round(val), "OK"
    except ValueError:
        return None, "WARN_WEIGHT"


def quality_level(flags: list[str]) -> str:
    if any(f.startswith("ERROR") for f in flags):
        return "ERROR"
    if any(f.startswith("WARN") for f in flags):
        return "WARN"
    return "OK"


# ── 主流程 ────────────────────────────────────────────────────────────────────

def run():
    src = BASE / "raw_data" / "shipping_2024_raw.xlsx"
    out = BASE / "output" / "cleaned_records.json"

    print(f"  讀取：{src}")
    wb = openpyxl.load_workbook(src)
    ws = wb["出貨明細"]

    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    idx = {h: i for i, h in enumerate(headers)}

    records = []
    warn_counts = {"WARN_DATE": 0, "WARN_VEHICLE": 0, "WARN_WEIGHT": 0}

    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_date    = row[idx["出貨日期"]]
        raw_vehicle = row[idx["車種（原始）"]]
        raw_weight  = row[idx["毛重(kg)_原始"]]

        date_iso, d_flag  = parse_date(raw_date)
        vehicle, v_flag   = normalize_vehicle(raw_vehicle)
        weight_kg, w_flag = normalize_weight(raw_weight)

        flags = []
        for flag in [d_flag, v_flag, w_flag]:
            if flag != "OK":
                flags.append(flag)
                if flag in warn_counts:
                    warn_counts[flag] += 1
        if not flags:
            flags = ["OK"]

        weight_tonne = round(weight_kg / 1000, 4) if weight_kg else None

        records.append({
            # 識別
            "id":           row[idx["運單號"]],
            # 清洗後
            "date":         date_iso,
            "dateRaw":      str(raw_date) if raw_date else "",
            "carrier":      row[idx["承運商"]] or "",
            "plate":        row[idx["車號"]] or "",
            "vehicle":      VEHICLE_CT[vehicle],
            "vehicleCode":  vehicle,
            "vehicleLabel": VEHICLE_LABELS[vehicle],
            "vehicleRaw":   str(raw_vehicle) if raw_vehicle else "",
            # 路線
            "origin":       row[idx["起點名稱"]] or "",
            "originAddr":   row[idx["起點地址"]] or "",
            "dest":         row[idx["迄點名稱"]] or "",
            "destAddr":     row[idx["迄點地址"]] or "",
            # 貨物
            "goodsType":    row[idx["貨物類型"]] or "",
            "pieces":       row[idx["件數"]] or 0,
            # 重量
            "weightKg":     weight_kg,
            "weight":       weight_tonne,   # 噸，供碳排計算用
            "weightRaw":    str(raw_weight) if raw_weight else "",
            # 其他
            "driverDist":   str(row[idx["司機自填距離(km)"]] or ""),
            "note":         str(row[idx["備註"]] or ""),
            # 品質
            "qualityFlags": flags,
            "qualityLevel": quality_level(flags),
        })

    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    ok_count   = sum(1 for r in records if r["qualityLevel"] == "OK")
    warn_count = sum(1 for r in records if r["qualityLevel"] == "WARN")
    err_count  = sum(1 for r in records if r["qualityLevel"] == "ERROR")

    print(f"  輸出：{out}")
    print(f"  總計：{len(records)} 筆")
    print(f"  品質：OK={ok_count}  WARN={warn_count}  ERROR={err_count}")
    for k, v in warn_counts.items():
        if v:
            print(f"    {k}: {v} 筆")


if __name__ == "__main__":
    run()
